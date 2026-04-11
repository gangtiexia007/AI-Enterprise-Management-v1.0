import json
from datetime import datetime, date
from io import BytesIO

from fastapi import APIRouter, Depends, UploadFile, HTTPException
from sqlalchemy import func, case
from sqlalchemy.orm import Session

from database import get_db
from models import PodOrder, PodProduct, Setting, AuditLog

router = APIRouter()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def extract_operator(store_name: str) -> str:
    if not store_name:
        return ""
    if "云往" in store_name or "新分销" in store_name:
        return ""
    parts = store_name.split("-")
    if len(parts) >= 3 and parts[-2].upper() in ("TK", "TMB", "SP"):
        return parts[-1]
    return ""


def get_niche_keyword_map(db: Session) -> dict:
    s = db.query(Setting).filter(Setting.key == "niche_keyword_map").first()
    if s and s.value:
        try:
            return json.loads(s.value)
        except json.JSONDecodeError:
            pass
    return {
        "潮牌街头": ["streetwear", "hip hop", "gangsta", "chicano", "swag", "urban"],
        "复古怀旧": ["vintage", "retro", "y2k", "90s", "80s", "old school", "nostalgia"],
        "宗教信仰": ["god", "jesus", "christian", "bible", "faith", "church", "blessed"],
        "情侣款": ["couple", "matching", "his and her", "boyfriend", "girlfriend"],
        "车迷机车": ["car", "motorcycle", "racing", "drift", "jdm", "porsche", "mustang", "truck"],
        "日系动漫": ["japanese", "anime", "manga", "otaku", "kawaii", "oriental", "samurai"],
        "运动健身": ["gym", "fitness", "sport", "basketball", "football", "workout", "boxing"],
        "宠物": ["cat", "dog", "pet", "kitten", "puppy", "paw", "shih tzu", "poodle"],
        "自然花卉": ["flower", "butterfly", "nature", "garden", "botanical", "sunflower"],
        "旅行城市": ["travel", "city", "country", "flag", "landmark", "world"],
        "骷髅摇滚": ["skull", "rock", "band", "metal", "punk", "gothic"],
        "咖啡生活": ["coffee", "latte", "cafe", "espresso", "tea"],
        "字母潮流": ["letter print", "alphabet", "monogram", "initial"],
    }


def classify_niche(title: str, keyword_map: dict) -> str:
    if not title:
        return ""
    title_lower = title.lower()
    for niche_name, keywords in keyword_map.items():
        for kw in keywords:
            if kw in title_lower:
                return niche_name
    return ""


def get_currency_rates(db: Session) -> dict:
    s = db.query(Setting).filter(Setting.key == "currency_rates").first()
    if s and s.value:
        try:
            return json.loads(s.value)
        except json.JSONDecodeError:
            pass
    return {"CNY": 1.0, "PHP": 0.125, "MYR": 1.55, "THB": 0.20}


STATUS_CATEGORY_MAP = {
    "已发货": "valid",
    "已完成": "valid",
    "确认收货": "valid",
    "待出库": "valid",
    "待揽收": "valid",
    "待处理": "valid",
    "已关闭": "cancelled",
    "安排失败": "failed",
    "售后中": "refund",
}

ORDER_COLUMN_MAP = {
    "ERP订单号": "erp_order_id",
    "平台订单号": "platform_order_id",
    "运单号": "tracking_number",
    "平台": "platform",
    "国家": "country",
    "店铺": "store",
    "平台产品ID": "platform_product_id",
    "平台产品SKU": "platform_sku",
    "平台产品规格": "product_spec",
    "元素SKU": "sku",
    "POD规格": "pod_spec",
    "产品数量": "quantity",
    "产品单价": "unit_price",
    "产品折扣价": "discount_price",
    "订单总金额": "total_amount",
    "实付金额": "paid_amount",
    "货币": "currency",
    "付款方式": "payment_method",
    "生产模式": "production_mode",
    "订单标签": "order_tag",
    "产品标题": "product_title",
    "付款时间": "payment_time",
    "运费": "shipping_fee",
    "预期发货时间": "expected_ship_time",
    "安排时间": "arrange_time",
    "创建时间": "created_time",
    "订单状态": "order_status",
    "失败原因": "failure_reason",
}

PRODUCT_HEADER_ALIASES = {
    "产品ID": "erp_product_id",
    "平台": "platform",
    "产品名称": "product_name",
    "商品名称": "product_name",
    "SKU编码": "platform_sku",
    "价格": "price",
    "店铺名称": "store",
    "店铺": "store",
    "创建时间": "upload_date",
}


# ---------------------------------------------------------------------------
# Row-parsing utilities
# ---------------------------------------------------------------------------

def _val(row, header_map: dict, field: str):
    idx = header_map.get(field)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _float(v):
    if v is None:
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _int(v):
    if v is None:
        return 1
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 1


def _str(v):
    if v is None:
        return ""
    return str(v).strip()


def _parse_datetime(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v
    try:
        return datetime.strptime(str(v).strip()[:19], "%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError):
        return None


def _parse_date(v):
    dt = _parse_datetime(v)
    return dt.date() if dt else None


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@router.post("/upload-orders")
async def upload_orders(file: UploadFile, db: Session = Depends(get_db)):
    import openpyxl

    content = await file.read()
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
    ws = wb["店铺订单"] if "店铺订单" in wb.sheetnames else wb.active

    header_map: dict[str, int] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            for col_idx, cell_val in enumerate(row):
                if cell_val and str(cell_val).strip() in ORDER_COLUMN_MAP:
                    header_map[ORDER_COLUMN_MAP[str(cell_val).strip()]] = col_idx
            break

    if not header_map or "erp_order_id" not in header_map:
        wb.close()
        raise HTTPException(status_code=400, detail="无法识别 Excel 列头，请确认文件格式")

    rates = get_currency_rates(db)
    niche_map = get_niche_keyword_map(db)

    existing_ids = set(r[0] for r in db.query(PodOrder.erp_order_id).all())

    new_count = 0
    skip_count = 0
    exclude_count = 0

    batch: list[PodOrder] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue

        erp_id = _str(_val(row, header_map, "erp_order_id"))
        if not erp_id:
            continue

        if erp_id in existing_ids:
            skip_count += 1
            continue

        store = _str(_val(row, header_map, "store"))
        operator = extract_operator(store)

        status_raw = _str(_val(row, header_map, "order_status"))
        status_cat = STATUS_CATEGORY_MAP.get(status_raw, "")

        currency = _str(_val(row, header_map, "currency")) or "PHP"
        paid = _float(_val(row, header_map, "paid_amount"))
        rate = rates.get(currency.upper(), rates.get("PHP", 0.125))
        paid_cny = round(paid * rate, 2)

        title = _str(_val(row, header_map, "product_title"))
        niche = classify_niche(title, niche_map)

        order = PodOrder(
            erp_order_id=erp_id,
            platform_order_id=_str(_val(row, header_map, "platform_order_id")),
            tracking_number=_str(_val(row, header_map, "tracking_number")),
            platform=_str(_val(row, header_map, "platform")),
            country=_str(_val(row, header_map, "country")),
            store=store,
            operator=operator,
            platform_product_id=_str(_val(row, header_map, "platform_product_id")),
            platform_sku=_str(_val(row, header_map, "platform_sku")),
            product_spec=_str(_val(row, header_map, "product_spec")),
            sku=_str(_val(row, header_map, "sku")),
            pod_spec=_str(_val(row, header_map, "pod_spec")),
            quantity=_int(_val(row, header_map, "quantity")),
            unit_price=_float(_val(row, header_map, "unit_price")),
            discount_price=_float(_val(row, header_map, "discount_price")),
            total_amount=_float(_val(row, header_map, "total_amount")),
            paid_amount=paid,
            paid_amount_cny=paid_cny,
            currency=currency,
            payment_method=_str(_val(row, header_map, "payment_method")),
            production_mode=_str(_val(row, header_map, "production_mode")),
            order_tag=_str(_val(row, header_map, "order_tag")),
            product_title=title,
            shipping_fee=_float(_val(row, header_map, "shipping_fee")),
            order_status=status_raw,
            status_category=status_cat,
            failure_reason=_str(_val(row, header_map, "failure_reason")),
            niche=niche,
            order_date=_parse_date(_val(row, header_map, "created_time")),
            payment_time=_parse_datetime(_val(row, header_map, "payment_time")),
            expected_ship_time=_parse_datetime(_val(row, header_map, "expected_ship_time")),
            arrange_time=_parse_datetime(_val(row, header_map, "arrange_time")),
            created_time=_parse_datetime(_val(row, header_map, "created_time")),
        )
        batch.append(order)
        existing_ids.add(erp_id)
        new_count += 1

        if len(batch) >= 500:
            db.bulk_save_objects(batch)
            db.commit()
            batch = []

    if batch:
        db.bulk_save_objects(batch)
        db.commit()

    wb.close()

    db.add(AuditLog(
        action="pod_order:import",
        detail=f"导入订单 Excel：新增 {new_count} 条，跳过 {skip_count} 条（重复），排除 {exclude_count} 条",
        actor="system",
        resource_type="pod_order",
    ))
    db.commit()

    return {
        "success": True,
        "new": new_count,
        "skipped": skip_count,
        "excluded": exclude_count,
        "total_in_db": db.query(PodOrder).count(),
    }


@router.post("/upload-products")
async def upload_products(file: UploadFile, db: Session = Depends(get_db)):
    import openpyxl

    content = await file.read()
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
    ws = wb["商品"] if "商品" in wb.sheetnames else wb.active

    header_map: dict[str, int] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            for col_idx, cell_val in enumerate(row):
                key = str(cell_val or "").strip()
                if key in PRODUCT_HEADER_ALIASES:
                    header_map[PRODUCT_HEADER_ALIASES[key]] = col_idx
            break

    niche_map = get_niche_keyword_map(db)

    existing = set(
        (r.erp_product_id, r.platform_sku)
        for r in db.query(PodProduct.erp_product_id, PodProduct.platform_sku).all()
    )

    new_count = 0
    skip_count = 0

    batch: list[PodProduct] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue

        pid = _str(_val(row, header_map, "erp_product_id"))
        psku = _str(_val(row, header_map, "platform_sku"))

        if (pid, psku) in existing:
            skip_count += 1
            continue

        store = _str(_val(row, header_map, "store"))
        operator = extract_operator(store)
        name = _str(_val(row, header_map, "product_name"))
        niche = classify_niche(name, niche_map)

        raw_date = _val(row, header_map, "upload_date")
        upload_dt: date | None = None
        if raw_date:
            if isinstance(raw_date, datetime):
                upload_dt = raw_date.date()
            else:
                try:
                    upload_dt = datetime.strptime(str(raw_date).strip()[:10], "%Y-%m-%d").date()
                except (ValueError, TypeError):
                    upload_dt = date.today()
        else:
            upload_dt = date.today()

        price_val = _float(_val(row, header_map, "price"))

        product = PodProduct(
            erp_product_id=pid,
            platform=_str(_val(row, header_map, "platform")),
            product_name=name,
            platform_sku=psku,
            price=price_val,
            store=store,
            operator=operator,
            niche=niche,
            upload_date=upload_dt,
        )
        batch.append(product)
        existing.add((pid, psku))
        new_count += 1

        if len(batch) >= 500:
            db.bulk_save_objects(batch)
            db.commit()
            batch = []

    if batch:
        db.bulk_save_objects(batch)
        db.commit()

    wb.close()

    db.add(AuditLog(
        action="pod_product:import",
        detail=f"导入产品 Excel：新增 {new_count} 条，跳过 {skip_count} 条（重复）",
        actor="system",
        resource_type="pod_product",
    ))
    db.commit()

    return {
        "success": True,
        "new": new_count,
        "skipped": skip_count,
        "total_in_db": db.query(PodProduct).count(),
    }


@router.get("/stats")
def pod_stats(db: Session = Depends(get_db)):
    total_orders = db.query(PodOrder).count()
    valid_orders = db.query(PodOrder).filter(PodOrder.status_category == "valid").count()
    total_products = db.query(PodProduct).count()
    operators = [
        r[0] for r in db.query(PodOrder.operator).filter(PodOrder.operator != "").distinct().all()
    ]
    return {
        "total_orders": total_orders,
        "valid_orders": valid_orders,
        "total_products": total_products,
        "operators": operators,
        "operator_count": len(operators),
    }


@router.get("/operator-kpi")
def operator_kpi(db: Session = Depends(get_db)):
    rows = (
        db.query(
            PodOrder.operator,
            func.sum(case((PodOrder.status_category == "valid", 1), else_=0)).label("valid_orders"),
            func.sum(case((PodOrder.status_category == "cancelled", 1), else_=0)).label("cancelled_orders"),
            func.sum(case((PodOrder.status_category == "refund", 1), else_=0)).label("refund_orders"),
            func.sum(case((PodOrder.status_category == "valid", PodOrder.paid_amount_cny), else_=0)).label("total_gmv_cny"),
            func.count(func.distinct(case((PodOrder.status_category == "valid", PodOrder.sku), else_=None))).label("unique_skus"),
        )
        .filter(PodOrder.operator != "")
        .group_by(PodOrder.operator)
        .all()
    )

    product_counts = dict(
        db.query(PodProduct.operator, func.count())
        .filter(PodProduct.operator != "")
        .group_by(PodProduct.operator)
        .all()
    )
    hit_counts = dict(
        db.query(PodProduct.operator, func.count())
        .filter(PodProduct.operator != "", PodProduct.has_order == 1)
        .group_by(PodProduct.operator)
        .all()
    )

    result = []
    for r in rows:
        valid = r.valid_orders or 0
        cancelled = r.cancelled_orders or 0
        total_base = valid + cancelled
        cancel_rate = round(cancelled / total_base * 100, 2) if total_base > 0 else 0

        total_products = product_counts.get(r.operator, 0)
        hit_products = hit_counts.get(r.operator, 0)
        hit_rate = round(hit_products / total_products * 100, 2) if total_products > 0 else 0

        result.append({
            "operator": r.operator,
            "valid_orders": valid,
            "cancelled_orders": cancelled,
            "refund_orders": r.refund_orders or 0,
            "total_gmv_cny": round(r.total_gmv_cny or 0, 2),
            "cancel_rate": cancel_rate,
            "unique_skus": r.unique_skus or 0,
            "total_products": total_products,
            "hit_products": hit_products,
            "hit_rate": hit_rate,
        })

    result.sort(key=lambda x: x["valid_orders"], reverse=True)
    return result


@router.get("/niche-stats")
def niche_stats(db: Session = Depends(get_db)):
    rows = (
        db.query(
            PodOrder.niche,
            func.sum(case((PodOrder.status_category == "valid", 1), else_=0)).label("valid_orders"),
            func.count(func.distinct(PodOrder.sku)).label("unique_skus"),
            func.count(func.distinct(PodOrder.operator)).label("operator_count"),
            func.sum(case((PodOrder.status_category == "valid", PodOrder.paid_amount_cny), else_=0)).label("total_gmv_cny"),
        )
        .filter(PodOrder.niche != "")
        .group_by(PodOrder.niche)
        .order_by(func.sum(case((PodOrder.status_category == "valid", 1), else_=0)).desc())
        .all()
    )
    return [
        {
            "niche": r.niche,
            "valid_orders": r.valid_orders or 0,
            "unique_skus": r.unique_skus or 0,
            "operator_count": r.operator_count or 0,
            "total_gmv_cny": round(r.total_gmv_cny or 0, 2),
        }
        for r in rows
    ]


@router.get("/sku-grading")
def sku_grading(db: Session = Depends(get_db)):
    valid_count = func.sum(case((PodOrder.status_category == "valid", 1), else_=0))
    rows = (
        db.query(
            PodOrder.sku,
            PodOrder.product_title,
            PodOrder.operator,
            PodOrder.niche,
            PodOrder.platform,
            valid_count.label("valid_orders"),
            func.sum(case((PodOrder.status_category == "valid", PodOrder.paid_amount_cny), else_=0)).label("gmv_cny"),
        )
        .filter(PodOrder.sku != "")
        .group_by(PodOrder.sku)
        .having(valid_count > 0)
        .order_by(valid_count.desc())
        .all()
    )

    result = []
    for r in rows:
        valid = r.valid_orders or 0
        if valid >= 20:
            grade = "S"
        elif valid >= 10:
            grade = "A"
        elif valid >= 5:
            grade = "B"
        elif valid >= 2:
            grade = "C"
        else:
            grade = "D"

        result.append({
            "sku": r.sku,
            "product_title": r.product_title or "",
            "operator": r.operator or "",
            "niche": r.niche or "",
            "platform": r.platform or "",
            "valid_orders": valid,
            "gmv_cny": round(r.gmv_cny or 0, 2),
            "grade": grade,
        })

    return result
